"""Gera uma prévia auditável de Codigos_Fanor2026.xlsx, sem escrever no banco.

Uso: python scripts/preview-codigos-fanor.py caminho/arquivo.xlsx --output caminho/previa.json
Requer openpyxl (`python -m pip install openpyxl`). O JSON de saída contém dados
do cliente; salve-o em `materiais-cliente/`, que está fora do repositório público.
"""

import argparse
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def label(value):
    return "" if value is None else str(value).strip()


def quantity(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not math.isfinite(value) or value <= 0:
        return None
    return value


def issue(issues, severity, kind, sheet, row, message):
    issues.append(
        {
            "severity": severity,
            "kind": kind,
            "sheet": sheet,
            "row": row,
            "message": message,
        }
    )


def preview(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    catalog_sheet = "XLS Milly"
    recipes_sheet = "RECETAS INTERMEDIOS TORTAS"
    missing = {catalog_sheet, recipes_sheet} - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"Abas ausentes: {', '.join(sorted(missing))}")

    issues = []
    catalog = []
    by_code = {}
    by_old_code = {}
    for row, cells in enumerate(workbook[catalog_sheet].iter_rows(values_only=True), 1):
        if row == 1:
            continue
        old_code, code, description, unit = (label(v) for v in cells[:4])
        if not any((old_code, code, description, unit)):
            continue
        if not code or not description or not unit:
            issue(issues, "blocker", "incomplete_catalog_item", catalog_sheet, row,
                  "Código, descrição e unidade são obrigatórios.")
            continue
        if code in by_code:
            issue(issues, "blocker", "duplicate_code", catalog_sheet, row,
                  f"Código {code} já aparece na linha {by_code[code]['row']}.")
        if old_code and old_code in by_old_code:
            issue(issues, "blocker", "duplicate_old_code", catalog_sheet, row,
                  f"Código antigo {old_code} já aparece na linha {by_old_code[old_code]['row']}.")
        item = {"row": row, "old_code": old_code or None, "code": code,
                "description": description, "unit": unit}
        catalog.append(item)
        by_code.setdefault(code, item)
        if old_code:
            by_old_code.setdefault(old_code, item)
        if "\ufffd" in description:
            issue(issues, "review", "damaged_text", catalog_sheet, row,
                  f"Descrição de {code} contém caractere de substituição.")

    recipes = []
    current = None
    for row, cells in enumerate(workbook[recipes_sheet].iter_rows(values_only=True), 1):
        if row == 1:
            continue
        output_code, yield_value, input_code, input_qty = (cells[i] if i < len(cells) else None for i in range(4))
        output_code, input_code = label(output_code), label(input_code)
        if not output_code and yield_value is None and not input_code and input_qty is None:
            continue
        if output_code:
            item = by_code.get(output_code)
            old_item = by_old_code.get(output_code) if item is None else None
            if old_item:
                issue(issues, "review", "old_code_in_recipe", recipes_sheet, row,
                      f"Saída {output_code} usa código antigo; código novo: {old_item['code']}.")
                item = old_item
            elif item is None:
                issue(issues, "blocker", "unknown_output_code", recipes_sheet, row,
                      f"Saída {output_code} não consta no catálogo.")
            if quantity(yield_value) is None:
                issue(issues, "blocker", "missing_or_invalid_yield", recipes_sheet, row,
                      f"Rendimento de {output_code} está vazio ou não é positivo.")
            current = {"row": row, "source_code": output_code,
                       "code": item["code"] if item else None,
                       "yield_quantity": quantity(yield_value),
                       "yield_unit": item["unit"] if item else None,
                       "ingredients": []}
            recipes.append(current)
        elif yield_value is not None:
            issue(issues, "blocker", "orphan_yield", recipes_sheet, row,
                  "Há rendimento sem código de saída.")
        if input_code or input_qty is not None:
            if current is None:
                issue(issues, "blocker", "orphan_ingredient", recipes_sheet, row,
                      "Ingrediente aparece antes da primeira fórmula.")
                continue
            item = by_code.get(input_code)
            old_item = by_old_code.get(input_code) if item is None else None
            if old_item:
                issue(issues, "review", "old_ingredient_code", recipes_sheet, row,
                      f"Insumo {input_code} usa código antigo; código novo: {old_item['code']}.")
                item = old_item
            elif item is None:
                issue(issues, "blocker", "unknown_ingredient_code", recipes_sheet, row,
                      f"Insumo {input_code or '(vazio)'} não consta no catálogo.")
            if quantity(input_qty) is None:
                issue(issues, "blocker", "missing_or_invalid_ingredient_quantity", recipes_sheet, row,
                      f"Quantidade de {input_code or '(vazio)'} está vazia ou não é positiva.")
            current["ingredients"].append(
                {"row": row, "source_code": input_code or None,
                 "code": item["code"] if item else None,
                 "quantity": quantity(input_qty),
                 "unit": item["unit"] if item else None}
            )

    seen_recipes = {}
    for recipe in recipes:
        code = recipe["code"]
        if code in seen_recipes:
            issue(issues, "blocker", "duplicate_recipe_output", recipes_sheet, recipe["row"],
                  f"Receita de {code} já aparece na linha {seen_recipes[code]}.")
        else:
            seen_recipes[code] = recipe["row"]
        if not recipe["ingredients"]:
            issue(issues, "blocker", "recipe_without_ingredients", recipes_sheet, recipe["row"],
                  f"Receita de {code} não tem insumos.")
        if code == "MPP-100100" and {"MPP-100CLA", "MPP-100YEM"}.issubset(
            {ingredient["code"] for ingredient in recipe["ingredients"]}
        ):
            issue(issues, "review", "egg_separation_direction", recipes_sheet, recipe["row"],
                  "A fórmula mostra ovo inteiro como saída de clara e gema; confirmar o sentido da transformação.")

    workbook.close()
    counts = Counter(item["severity"] for item in issues)
    return {
        "source_file": path.name,
        "summary": {
            "catalog_items": len(catalog),
            "items_with_old_code": sum(bool(item["old_code"]) for item in catalog),
            "recipes": len(recipes),
            "ingredient_lines": sum(len(recipe["ingredients"]) for recipe in recipes),
            "blockers": counts["blocker"],
            "reviews": counts["review"],
        },
        "issues": issues,
        "catalog": catalog,
        "recipes": recipes,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = preview(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result["summary"], ensure_ascii=False))
    print(f"Prévia: {args.output}")


if __name__ == "__main__":
    main()
