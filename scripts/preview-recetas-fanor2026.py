"""Valida o catálogo e as fórmulas de Códigos _ Recetas_Fanor2026.xlsx.

Gera uma prévia JSON para o painel, sem criar produtos, estoque ou custos.
Uso: python scripts/preview-recetas-fanor2026.py arquivo.xlsx --output previa.json
"""

import argparse
import hashlib
import json
import math
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


def text(value):
    return "" if value is None else str(value).strip()


def positive(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value > 0


def issue(issues, severity, kind, sheet, row, message):
    issues.append({"severity": severity, "kind": kind, "sheet": sheet, "row": row, "message": message})


def converted_amount(amount, factor):
    result = Decimal(str(amount)) * Decimal(str(factor))
    return int(result) if result == result.to_integral_value() else float(result)


def build_preview(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if not {"Códigos", "RECETAS"}.issubset(workbook.sheetnames):
        raise ValueError("A planilha precisa conter as abas Códigos e RECETAS.")

    issues = []
    catalog = []
    by_code = {}
    for row, cells in enumerate(workbook["Códigos"].iter_rows(values_only=True), 1):
        if row == 1:
            if [text(v) for v in cells[:6]] != ["COD.", "TIPO", "DESCR.", "UNID. STOCK", "FACTOR CONV.", "UNID. RECETA"]:
                raise ValueError("Cabeçalhos da aba Códigos não correspondem ao formato esperado.")
            continue
        code, item_type, description, stock_unit, factor, recipe_unit, cost, minimum_lot = (list(cells[:8]) + [None] * 8)[:8]
        code, item_type, description, stock_unit, recipe_unit = map(text, (code, item_type, description, stock_unit, recipe_unit))
        if not any((code, item_type, description, stock_unit, factor, recipe_unit, cost, minimum_lot)):
            continue
        if not code or not description or not stock_unit or not recipe_unit:
            issue(issues, "blocker", "incomplete_catalog_item", "Códigos", row,
                  "Código, descripción y unidades son obligatorios.")
        if code in by_code:
            issue(issues, "blocker", "duplicate_code", "Códigos", row,
                  f"El código {code} ya aparece en la fila {by_code[code]['row']}.")
        if not positive(factor):
            issue(issues, "blocker", "invalid_conversion_factor", "Códigos", row,
                  f"El factor de conversión de {code} debe ser positivo.")
        elif stock_unit == recipe_unit and factor != 1:
            issue(issues, "review", "same_unit_conversion", "Códigos", row,
                  f"{code} tiene unidad {stock_unit} en ambos lados, pero factor {factor:g}; confirmar el sentido.")
        if "\ufffd" in description:
            issue(issues, "review", "damaged_text", "Códigos", row,
                  f"La descripción de {code} contiene un carácter dañado.")
        item = {
            "row": row, "old_code": None, "code": code, "description": description,
            "unit": recipe_unit, "type": item_type, "stock_unit": stock_unit,
            "conversion_factor": factor if positive(factor) else None,
            "cost": cost if isinstance(cost, (int, float)) else None,
            "minimum_lot": minimum_lot if isinstance(minimum_lot, (int, float)) else None,
        }
        catalog.append(item)
        if code:
            by_code.setdefault(code, item)

    recipes = []
    current = None
    conversion_count = 0
    for row, cells in enumerate(workbook["RECETAS"].iter_rows(values_only=True), 1):
        if row == 1:
            if [text(v) for v in cells[:4]] != ["CODIGO", "RENDIMIENTO", "CODIGO", "CANTIDAD"]:
                raise ValueError("Cabeçalhos da aba RECETAS não correspondem ao formato esperado.")
            continue
        output_code, yield_value, input_code, input_qty, weight, source_unit, calculation = (list(cells[:7]) + [None] * 7)[:7]
        output_code, input_code, source_unit = map(text, (output_code, input_code, source_unit))
        if not any((output_code, yield_value, input_code, input_qty, weight, source_unit, calculation)):
            continue
        if output_code:
            item = by_code.get(output_code)
            if item is None:
                issue(issues, "blocker", "unknown_output_code", "RECETAS", row,
                      f"La salida {output_code} no figura en Códigos.")
            if not positive(yield_value):
                issue(issues, "blocker", "missing_or_invalid_yield", "RECETAS", row,
                      f"El rendimiento de {output_code} falta o no es positivo.")
            current = {
                "row": row, "source_code": output_code, "code": output_code if item else None,
                "yield_quantity": yield_value if positive(yield_value) else None,
                "yield_unit": item["unit"] if item else None,
                "source_weight": weight, "ingredients": [],
            }
            recipes.append(current)
        elif yield_value is not None:
            issue(issues, "blocker", "orphan_yield", "RECETAS", row,
                  "Hay un rendimiento sin código de salida.")

        if input_code or input_qty is not None:
            if current is None:
                issue(issues, "blocker", "orphan_ingredient", "RECETAS", row,
                      "Hay un ingrediente antes de la primera fórmula.")
                continue
            item = by_code.get(input_code)
            if item is None:
                issue(issues, "blocker", "unknown_ingredient_code", "RECETAS", row,
                      f"El ingrediente {input_code or '(vacío)'} no figura en Códigos.")
            if not positive(input_qty):
                issue(issues, "blocker", "missing_or_invalid_ingredient_quantity", "RECETAS", row,
                      f"La cantidad de {input_code or '(vacío)'} falta o no es positiva.")
            quantity, unit, factor_used = input_qty if positive(input_qty) else None, item["unit"] if item else source_unit or None, None
            if item and positive(input_qty):
                if source_unit == item["unit"]:
                    pass
                elif source_unit == item["stock_unit"] and positive(item["conversion_factor"]):
                    factor_used = item["conversion_factor"]
                    quantity = converted_amount(input_qty, factor_used)
                    conversion_count += 1
                else:
                    quantity = None
                    issue(issues, "blocker", "incompatible_ingredient_unit", "RECETAS", row,
                          f"{input_code}: {source_unit or '(vacío)'} no coincide con {item['stock_unit']} ni {item['unit']}.")
            current["ingredients"].append({
                "row": row, "source_code": input_code or None,
                "code": input_code if item else None, "quantity": quantity, "unit": unit,
                "source_quantity": input_qty if positive(input_qty) else None,
                "source_unit": source_unit or None, "conversion_factor": factor_used,
                "source_calculation": calculation,
            })

    seen = {}
    for recipe in recipes:
        if recipe["source_code"] in seen:
            issue(issues, "blocker", "duplicate_recipe_output", "RECETAS", recipe["row"],
                  f"La fórmula {recipe['source_code']} ya aparece en la fila {seen[recipe['source_code']]}.")
        seen.setdefault(recipe["source_code"], recipe["row"])
        if not recipe["ingredients"]:
            issue(issues, "blocker", "recipe_without_ingredients", "RECETAS", recipe["row"],
                  f"La fórmula {recipe['source_code']} no tiene ingredientes.")

    # A clara e a gema aparecem como duas saídas separadas do mesmo ovo; executar
    # as duas receitas como processos independentes poderia baixar o ovo duas vezes.
    for code in ("MPP-100CLA", "MPP-100YEM"):
        recipe = next((recipe for recipe in recipes if recipe["source_code"] == code), None)
        if recipe and any(ingredient["source_code"] == "MPP-100100" for ingredient in recipe["ingredients"]):
            issue(issues, "review", "egg_separation_multi_output", "RECETAS", recipe["row"],
                  "Clara y yema son dos salidas del mismo huevo; confirmar un único proceso con dos resultados.")

    workbook.close()
    counts = Counter(entry["severity"] for entry in issues)
    return {
        "source_file": path.name,
        "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "recipe_sheet": "RECETAS",
        "summary": {
            "catalog_items": len(catalog), "items_with_old_code": 0,
            "recipes": len(recipes),
            "ingredient_lines": sum(len(recipe["ingredients"]) for recipe in recipes),
            "converted_ingredient_lines": conversion_count,
            "items_without_cost": sum(item["cost"] is None for item in catalog),
            "blockers": counts["blocker"], "reviews": counts["review"],
        },
        "issues": issues, "catalog": catalog, "recipes": recipes,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    preview = build_preview(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(preview, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(preview["summary"], ensure_ascii=False))
    print(f"Prévia: {args.output}")


if __name__ == "__main__":
    main()
